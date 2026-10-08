import xml.etree.ElementTree as ET
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
from tkinter import Tk, filedialog, messagebox


ARQUIVO_EXCEL = "teste.xlsx"

LINHA_INICIAL = 6
LINHA_FINAL = 110

MESES = {
    1: "Mês 01",
    2: "Mês 02",
    3: "Mês 03",
    4: "Mês 04",
}

NS = {
    "nfe": "http://www.portalfiscal.inf.br/nfe"
}


def selecionar_xmls():

    root = Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    arquivos = filedialog.askopenfilenames(

        title="Selecione os XMLs",

        filetypes=[
            ("Arquivos XML", "*.xml")
        ]

    )

    root.destroy()

    return [Path(a) for a in arquivos]


def texto(elemento):
    if elemento is None:
        return ""
    return elemento.text.strip()


def ler_xml(arquivo):

    tree = ET.parse(arquivo)
    root = tree.getroot()

    infNFe = root.find(".//nfe:infNFe", NS)

    if infNFe is None:
        raise Exception("infNFe não encontrada.")

    cliente = texto(
        infNFe.find("./nfe:dest/nfe:xNome", NS)
    )

    numero_nf = texto(
        infNFe.find("./nfe:ide/nfe:nNF", NS)
    )

    emissao = texto(
        infNFe.find("./nfe:ide/nfe:dhEmi", NS)
    )

    emissao = datetime.fromisoformat(
        emissao.replace("Z", "+00:00")
    ).date()

    duplicatas = []

    cobr = infNFe.find("./nfe:cobr", NS)

    if cobr is not None:

        for dup in cobr.findall("./nfe:dup", NS):

            numero = texto(
                dup.find("./nfe:nDup", NS)
            )

            vencimento = texto(
                dup.find("./nfe:dVenc", NS)
            )

            valor = texto(
                dup.find("./nfe:vDup", NS)
            )

            duplicatas.append({

                "numero": numero,

                "vencimento":
                    datetime.strptime(
                        vencimento,
                        "%Y-%m-%d"
                    ).date(),

                "valor":
                    float(valor),

                "avista": False

            })

    if len(duplicatas) == 0:

        total = texto(
            infNFe.find(
                "./nfe:total/nfe:ICMSTot/nfe:vNF",
                NS
            )
        )

        duplicatas.append({

            "numero": "",

            "vencimento": emissao,

            "valor": float(total),

            "avista": True

        })

    return {

        "cliente": cliente,

        "nota": numero_nf,

        "emissao": emissao,

        "parcelas": duplicatas

    }


def encontrar_aba(workbook, mes):

    nome = MESES[mes]

    if nome not in workbook.sheetnames:
        raise Exception(f"A aba '{nome}' não existe.")

    return workbook[nome]


def documento_existe(ws, documento):

    for linha in range(LINHA_INICIAL, LINHA_FINAL + 1):

        doc = ws.cell(linha, 3).value

        if doc is None:
            continue

        if str(doc).strip() == documento:
            return True

    return False


def primeira_linha_livre(ws):

    for linha in range(LINHA_INICIAL, LINHA_FINAL + 1):

        if ws.cell(linha, 1).value is None:
            return linha

    raise Exception("Não há espaço livre na planilha.")


def ordenar_aba(ws):

    registros = []

    # Lê somente a área da tabela
    for linha in range(LINHA_INICIAL, LINHA_FINAL + 1):

        if ws.cell(linha, 1).value is None:
            continue

        registros.append([
            ws.cell(linha, 1).value,  # Vencimento
            ws.cell(linha, 2).value,  # Cliente
            ws.cell(linha, 3).value,  # Documento
            ws.cell(linha, 4).value,  # Valor
            ws.cell(linha, 5).value,  # Emissão
            ws.cell(linha, 6).value,  # Recebimento
            ws.cell(linha, 7).value   # Valor Recebido
        ])

    # Ordena pela data de vencimento
    registros.sort(
        key=lambda x: x[0].date() if isinstance(x[0], datetime) else x[0]
    )

    # Limpa somente a área da tabela
    for linha in range(LINHA_INICIAL, LINHA_FINAL + 1):

        for coluna in range(1, 8):
            ws.cell(linha, coluna).value = None

    # Escreve novamente
    linha = LINHA_INICIAL

    for registro in registros:

        for coluna, valor in enumerate(registro, start=1):
            ws.cell(linha, coluna).value = valor

        linha += 1

    # Reaplica os formatos de data e moeda
    for linha in range(LINHA_INICIAL, linha):

        ws.cell(linha, 1).number_format = "dd/mm/yyyy"
        ws.cell(linha, 4).number_format = 'R$ #,##0.00'
        ws.cell(linha, 5).number_format = "dd/mm/yyyy"
        ws.cell(linha, 7).number_format = 'R$ #,##0.00'


def inserir_xml(workbook, dados):

    abas_modificadas = set()

    for parcela in dados["parcelas"]:

        vencimento = parcela["vencimento"]


        aba = encontrar_aba(
            workbook,
            vencimento.month
        )


        if parcela["avista"]:

            documento = dados["nota"]

        else:

            documento = f'{dados["nota"]}/{parcela["numero"]}'


        if documento_existe(aba, documento):

            print(f"{documento} já existe. Ignorado.")

            continue


        linha = primeira_linha_livre(aba)


        aba.cell(linha, 1).value = vencimento

        aba.cell(linha, 2).value = dados["cliente"]

        aba.cell(linha, 3).value = documento

        aba.cell(linha, 4).value = parcela["valor"]

        aba.cell(linha, 5).value = dados["emissao"]


        aba.cell(linha, 1).number_format = "dd/mm/yyyy"

        aba.cell(linha, 4).number_format = 'R$ #,##0.00'

        aba.cell(linha, 5).number_format = "dd/mm/yyyy"


        print(f"Adicionado: {documento}")

        abas_modificadas.add(aba)

    # Ordena apenas as abas que foram modificadas
    for aba in abas_modificadas:
        ordenar_aba(aba)

def main():

    if not Path(ARQUIVO_EXCEL).exists():
        print(f"Arquivo '{ARQUIVO_EXCEL}' não encontrado.")
        return

    arquivos_xml = selecionar_xmls()

    if len(arquivos_xml) == 0:
        print("Nenhum XML selecionado.")
        return

    print(f"Foram encontrados {len(arquivos_xml)} XML(s).\n")

    workbook = load_workbook(
        ARQUIVO_EXCEL,
        keep_links=False
    )


    importados = 0
    erros = 0

    for arquivo in arquivos_xml:

        print("=" * 70)
        print(f"Processando: {arquivo.name}")

        try:

            dados = ler_xml(arquivo)

            print(f"Cliente : {dados['cliente']}")
            print(f"NF      : {dados['nota']}")
            print(f"Parcelas: {len(dados['parcelas'])}")

            inserir_xml(workbook, dados)

            importados += 1

        except Exception as e:

            erros += 1
            print(f"ERRO: {e}")

    workbook.save("teste (feito).xlsx")

    print("\n" + "=" * 70)
    print("FINALIZADO")
    print(f"XMLs processados : {len(arquivos_xml)}")
    print(f"Importados       : {importados}")
    print(f"Erros            : {erros}")

    messagebox.showinfo(
        "Finalizado",
        f"Importação concluída!\n\nXMLs processados: {len(arquivos_xml)}\nErros: {erros}"
    )


if __name__ == "__main__":
    main()